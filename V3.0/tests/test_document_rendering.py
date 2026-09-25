"""CHG-026 — Professional Document Engine tests (offline, deterministic).

Prove, WITHOUT any AI call:
- the structured document model validates (schema gate);
- the builder maps generated content + profile FACTS and invents nothing;
- the DOCX renderer produces A4 documents with real structures (list numbering,
  w:hyperlink, PAGE/NUMPAGES footer fields, keep-with-next pagination guards);
- content integrity: every approved source token survives into the rendered
  document, and markup (Markdown/JSON) never leaks into the final document;
- a rendering failure is an EXCEPTION, never a silent content change;
- LibreOffice PDF conversion is probed: skipped with a clear reason when
  LibreOffice is absent, fully verified when it is present.
"""

from __future__ import annotations

import io
import re
import zipfile

import pytest

from documents.builders import build_cover_letter_document, build_cv_document
from documents.docx_renderer import DocumentRenderError, render_document
from documents.models import DocumentModel
from documents.pdf_renderer import PdfConversionError, convert_docx_to_pdf, libreoffice_available
from documents.validation import (
    check_content_integrity,
    model_source_texts,
    validate_document_model,
    validate_rendered_document,
)

PROFILE = {
    "identity": {
        "full_name": "Aymen Abdellaoui",
        "email": "aymen.a@hotmail.com",
        "phone": "+216 22142743",
        "linkedin": "linkedin.com/in/aymen-it-pm",
        "location": "Tunisia",
        "whatsapp": "+216 22 142 743",
        "has_driver_license": True,
    },
    "experience": [
        {"title": "Program & Program Manager", "company": "IPACT",
         "location": "Tunis, Tunisia", "start_date": "01.2022", "end_date": "Present"},
    ],
}

GENERATED_CV = {
    "summary": ("ERP & IT Programme Manager with 16+ years of experience leading enterprise "
                "digital transformation, custom ERP development, and enterprise software "
                "delivery across manufacturing and service industries."),
    "experience": [
        {
            "title": "Program Manager - SaaS & Digital Transformation",
            "company": "IPACT (Tunis, Tunisia)",
            "start_date": "01.2022",
            "end_date": "Present",
            "description": ("Led digital transformation programs delivering SaaS platforms and "
                            "enterprise applications through hybrid PMO governance."),
            "achievements": [
                "**ClicToEat Multi-Tenant Restaurant SaaS platform** : Supported 20+ restaurants "
                "across 4 countries.",
                "Delivery Approach: Agile/Scrum",
            ],
        },
    ],
    "skills": {"technical": ["Python", "SQL"], "methodologies": ["Agile"], "tools": ["Docker"]},
    "education": [{"degree": "Master in IT Project Management", "institution": "ESSEC",
                   "year": "2012"}],
    "certifications": [{"name": "PMP", "issuer": "PMI", "year": "2020"}],
    "languages": [{"language": "French", "level": "Native"},
                  {"language": "English", "level": "C1"}],
    "full_text": "ERP & IT Programme Manager with 16+ years of experience.",
}

GENERATED_LETTER = {
    "greeting": "Dear Mr. Tony Yong,",
    "opening": ("I am applying for the SAP PM Consultant position at Enexis with hands-on "
                "experience leading ERP-driven digital transformation projects."),
    "body": ("In my current role, I lead ERP implementation projects involving stakeholder "
             "workshops, functional specification development, and system validation (UAT)."),
    "closing": "Kind regards,",
    "full_text": "",
}

JOB = {"job_id": "34146NL", "title": "ERP Implementation Project Manager",
       "job_json": {"job_title": "ERP Implementation Project Manager", "country": "Netherlands"}}


@pytest.fixture()
def cv_model():
    return build_cv_document(GENERATED_CV, PROFILE, meta={"job_id": JOB["job_id"]})


@pytest.fixture()
def letter_model():
    return build_cover_letter_document(GENERATED_LETTER, PROFILE, JOB)


@pytest.fixture()
def cv_docx(cv_model):
    return render_document(cv_model)
# ---------------------------------------------------------------- schema gate
def test_cv_model_passes_schema(cv_model):
    report = validate_document_model(cv_model)
    assert report["ok"] is True, report["errors"]


def test_letter_model_passes_schema(letter_model):
    report = validate_document_model(letter_model)
    assert report["ok"] is True, report["errors"]


def test_model_without_name_fails_schema():
    model = build_cv_document(GENERATED_CV, {"identity": {}, "experience": []})
    model.candidate.name = ""
    report = validate_document_model(model)
    assert report["ok"] is False
    assert any("name" in error for error in report["errors"])


# ------------------------------------------------------------------- builder
def test_builder_maps_profile_facts_and_omits_missing(cv_model):
    header = cv_model.candidate
    assert header.name == "Aymen Abdellaoui"
    # professional title = FACT: the most recent profile experience title
    assert header.professional_title == "Program & Program Manager"
    assert header.status_line == "Tunisia | Driving Licence (B)"
    texts = " | ".join(
        item.text if hasattr(item, "text") else str(item) for item in header.contact_items
    )
    assert "aymen.a@hotmail.com" in texts and "+216 22142743" in texts
    assert "linkedin.com/in/aymen-it-pm" in texts


def test_builder_uses_jd_facts_for_subject_and_destination(letter_model):
    assert letter_model.subject == (
        "Application – ERP Implementation Project Manager (Job ID 34146NL)"
    )
    assert letter_model.location_line == "Netherlands"
    assert letter_model.salutation == "Dear Mr. Tony Yong,"
    assert letter_model.closing == "Kind regards,"
    assert letter_model.signature == "Aymen Abdellaoui"


def test_builder_does_not_invent_missing_facts():
    profile = {"identity": {"full_name": "Jane Doe"}, "experience": []}
    model = build_cv_document({"summary": "Senior engineer."}, profile)
    assert model.candidate.professional_title == ""
    assert model.candidate.status_line == ""
    assert model.candidate.contact_items == []
# ------------------------------------------------------------------- DOCX QA
def test_cv_docx_uses_real_structures(cv_docx):
    archive = zipfile.ZipFile(io.BytesIO(cv_docx))
    doc = archive.read("word/document.xml").decode("utf-8")
    # A4 (11906 x 16838 twips) + the measured margins
    assert 'w:w="11906"' in doc and 'w:h="16838"' in doc
    assert 'w:left="1440"' in doc  # 2.54 cm
    # TRUE list structures (not typed hyphens) with the "-" glyph
    assert doc.count("<w:numPr") >= 3
    assert '<w:lvlText w:val="-"' in archive.read("word/numbering.xml").decode("utf-8")
    # real hyperlinks with the controlled (non-default) colour
    assert doc.count("<w:hyperlink") >= 2
    assert 'w:val="1155CC"' in doc
    # pagination hygiene: headings/entries kept with their content
    assert "<w:keepNext/>" in doc
    # footer with automatic page numbering fields
    footers = "".join(
        archive.read(name).decode("utf-8")
        for name in archive.namelist() if "footer" in name and name.endswith(".xml")
    )
    assert "PAGE" in footers and "NUMPAGES" in footers
    # font family of the reference baseline
    assert "Calibri" in doc


def test_letter_docx_renders_full_hierarchy(letter_model):
    import docx as docx_module

    document = docx_module.Document(io.BytesIO(render_document(letter_model)))
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    assert "Aymen Abdellaoui" in text
    assert "aymen.a@hotmail.com" in text
    assert "LinkedIn:" in text
    assert re.search(r"\d{2} \w+ \d{4}", text)  # date line
    assert text.startswith("Aymen Abdellaoui")  # contact block leads the letter
    assert "Subject : Application – ERP Implementation Project Manager (Job ID 34146NL)" in text
    assert "Dear Mr. Tony Yong," in text
    assert "Kind regards," in text
    # the generated wording survives verbatim
    assert "ERP implementation projects involving stakeholder" in text


def test_rendered_document_passes_integrity(cv_model, cv_docx):
    report = validate_rendered_document(cv_model, cv_docx)
    assert report["ok"] is True, report["errors"]
    assert report["integrity"]["missing_tokens"] == []
    assert report["integrity"]["markup_leaks"] == []


def test_rendered_letter_passes_integrity(letter_model):
    report = validate_rendered_document(letter_model, render_document(letter_model))
    assert report["ok"] is True, report["errors"]


# ------------------------------------------------------- content preservation
def test_integrity_detects_content_loss():
    report = check_content_integrity(
        ["Led ERP programs across Germany and Tunisia"], "Led ERP programs across Germany"
    )
    assert report["ok"] is False
    assert "tunisia" in report["missing_tokens"]


def test_integrity_detects_markup_leak():
    report = check_content_integrity(["Led ERP programs"], "Led **ERP** programs")
    assert report["ok"] is False
    assert report["markup_leaks"]


def test_inline_markdown_becomes_real_bold_runs(cv_model):
    """'**ClicToEat...**' becomes a bold run - the markers must NOT survive."""
    import docx as docx_module

    document = docx_module.Document(io.BytesIO(render_document(cv_model)))
    all_runs = [run for paragraph in document.paragraphs for run in paragraph.runs]
    assert not any("**" in run.text for run in all_runs)
    clic_runs = [run for run in all_runs if "ClicToEat" in run.text]
    assert clic_runs and all(run.bold for run in clic_runs)


def test_unknown_doc_type_is_a_clean_render_error():
    model = DocumentModel.model_validate({"doc_type": "unknown", "candidate": {"name": "Jane Doe"}})
    with pytest.raises(DocumentRenderError):
        render_document(model)


# ----------------------------------------------------------------- PDF (opt-in)
def test_pdf_conversion_without_libreoffice_reports_cleanly(tmp_path):
    from documents import pdf_renderer

    if libreoffice_available():
        pytest.skip("LibreOffice IS installed - the 'missing' path is not reachable")
    source = tmp_path / "cv.docx"
    source.write_bytes(render_document(build_cv_document(GENERATED_CV, PROFILE)))
    with pytest.raises(PdfConversionError) as excinfo:
        pdf_renderer.convert_docx_to_pdf(source)
    assert "LibreOffice" in str(excinfo.value)


@pytest.mark.skipif(not libreoffice_available(), reason="LibreOffice not installed on this host")
def test_pdf_conversion_produces_a4_with_full_content(cv_model, cv_docx, tmp_path):
    import fitz  # pymupdf

    source = tmp_path / "cv.docx"
    source.write_bytes(cv_docx)
    pdf_path = convert_docx_to_pdf(source)
    assert pdf_path.exists() and pdf_path.stat().st_size > 1000
    document = fitz.open(str(pdf_path))
    try:
        page = document[0]
        assert abs(page.rect.width - 595) < 3 and abs(page.rect.height - 842) < 3  # A4 pt
        text = "\n".join(page.get_text() for page in document)
        assert "Aymen Abdellaoui" in text
        integrity = check_content_integrity(
            model_source_texts(cv_model), text, footer_name=cv_model.candidate.name
        )
        assert integrity["ok"] is True, integrity["missing_tokens"]
    finally:
        document.close()


