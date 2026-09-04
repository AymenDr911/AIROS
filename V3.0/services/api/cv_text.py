"""AIROS V3 - CV text extraction (Slice 6, V2 parity).

Verbatim port of V2 ``services/cv_extractor.py`` section 2 (PDF + DOCX text
extraction), adapted from Streamlit ``UploadedFile`` objects to plain
``(filename, bytes)`` pairs:

- PDF: PyMuPDF preferred -> pypdf fallback (same order as V2)
- DOCX: python-docx paragraphs + tables + headers/footers (same as V2)
- multi-CV concat with ``=== CV: <name> ===`` markers (same as V2)

No AI logic here - extraction only; the gateway (ai/gateway.py) consumes the
text. This module never touches secrets (ISS-004).
"""

from __future__ import annotations

from io import BytesIO
from typing import List, Tuple


class CvTextError(Exception):
    """A CV could not be read/parsed (user-displayable)."""


def extract_text_from_file(filename: str, data: bytes) -> str:
    """Return the raw text of one CV file ('' when nothing could be read)."""
    name = (filename or "").lower()
    try:
        if name.endswith(".pdf"):
            return _extract_pdf_text(data)
        elif name.endswith(".docx"):
            return _extract_docx_text(data)
        else:
            return ""
    except CvTextError:
        raise
    except Exception as e:  # noqa: BLE001 - parity with V2 per-file catch
        raise CvTextError(f"Could not read {filename}: {e}") from e


def _extract_pdf_text(data: bytes) -> str:
    """Extract text with PyMuPDF (preferred) -> fallback to pypdf (V2 order)."""
    try:
        import pymupdf  # new recommended import (V2 parity)

        doc = pymupdf.open(stream=data, filetype="pdf")
        text_parts = []
        for page in doc:
            text = page.get_text("text")
            if text.strip():
                text_parts.append(text)
        doc.close()
        if text_parts:
            return "\n".join(text_parts).strip()
    except Exception as e:  # noqa: BLE001 - V2 falls back to pypdf on any error
        # V2 warned and fell through to pypdf; keep going the same way.
        _ = e

    # Fallback to pypdf
    from pypdf import PdfReader

    reader = PdfReader(BytesIO(data))
    text = ""
    for page in reader.pages:
        page_text = page.extract_text()
        if page_text:
            text += page_text + "\n"
    return text.strip()


def _extract_docx_text(data: bytes) -> str:
    from docx import Document

    doc = Document(BytesIO(data))
    parts = []

    for p in doc.paragraphs:
        if p.text.strip():
            parts.append(p.text)

    for table in doc.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                parts.append(" | ".join(cells))

    for section in doc.sections:
        for hdr in (section.header, section.first_page_header):
            if hdr is not None:
                for p in hdr.paragraphs:
                    if p.text.strip():
                        parts.append(p.text)
        for ftr in (section.footer, section.first_page_footer):
            if ftr is not None:
                for p in ftr.paragraphs:
                    if p.text.strip():
                        parts.append(p.text)

    return "\n".join(parts).strip()


def extract_text_from_multiple(files: List[Tuple[str, bytes]]) -> str:
    """Concatenate the text of several CVs (V2 marker format)."""
    all_text = []
    for name, data in files:
        text = extract_text_from_file(name, data)
        if text:
            all_text.append(f"=== CV: {name} ===\n{text}")
    return "\n\n".join(all_text)
