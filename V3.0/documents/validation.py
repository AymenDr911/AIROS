"""AIROS Document QA (CHG-026): schema validation + content integrity.

Two independent gates, both deterministic (no AI call):

1. Schema validation  - the structured model is complete enough to render
   (required fields per document type; no empty content sections).
2. Content integrity  - the rendered document contains ALL of the approved
   source content: the normalized source tokens must be a subset of the
   normalized extracted tokens (missing = LOSS), and no raw Markdown / JSON /
   HTML markup may leak into the final document (guide section 10). The
   renderer may ADD layout artifacts (footer name, "Page : X/Y"); it may never
   REMOVE content. Overflow is REPORTED, never solved by deleting content.
"""

from __future__ import annotations

import re
from typing import Any, Dict, List, Optional, Sequence

from documents.models import DocumentModel, Link

# tokens = words/abbreviations; keep '+', '#', '.' inside tokens (C#, .NET...)
_TOKEN = re.compile(r"[A-Za-z][A-Za-z0-9\+\#\.\-]*")
_MARKUP_LEAK = re.compile(r"\*\*|``|~~|\{\s*\"|</?[a-z]+[ >]")
# layout artifacts the renderer is ALLOWED to add (never counted as additions)
_ALLOWED_EXTRA = {"page"}


def _tokens(text: str) -> List[str]:
    return [token.lower() for token in _TOKEN.findall(text or "")]


def _extract_docx_text(docx_bytes: bytes) -> str:
    """All paragraph text of a DOCX (python-docx round-trip, no external tool)."""
    import io

    from docx import Document

    document = Document(io.BytesIO(docx_bytes))
    parts: List[str] = []
    for paragraph in document.paragraphs:
        if paragraph.text.strip():
            parts.append(paragraph.text)
    for section in document.sections:
        for paragraph in section.footer.paragraphs:
            if paragraph.text.strip():
                parts.append(paragraph.text)
    return "\n".join(parts)


def _extract_pdf_text(pdf_path) -> str:
    """All text of a PDF (pymupdf; V2-parity extraction dependency)."""
    import fitz  # pymupdf

    document = fitz.open(str(pdf_path))
    try:
        return "\n".join(page.get_text() for page in document)
    finally:
        document.close()


def validate_document_model(model: DocumentModel) -> Dict[str, Any]:
    """Schema gate: required structure per document type. Never mutates data."""
    errors: List[str] = []
    if model.doc_type not in ("cv", "cover_letter"):
        errors.append(f"unknown doc_type {model.doc_type!r}")
    if not model.candidate.name.strip():
        errors.append("candidate.name is required")
    if model.doc_type == "cv":
        if not model.sections:
            errors.append("a CV needs at least one section")
        for section in model.sections:
            has_content = bool(
                section.paragraphs or section.bullets or section.experience
                or section.education or section.entries
            )
            if not section.heading.strip() or not has_content:
                errors.append(f"section {section.heading!r} is empty or unnamed")
    else:
        if not (model.salutation.strip() or model.sections):
            errors.append("a cover letter needs a salutation or body paragraphs")
        body_paragraphs = [p for s in model.sections for p in s.paragraphs]
        if not any(p.text.strip() for p in body_paragraphs):
            errors.append("cover letter body is empty")
        if not model.closing.strip():
            errors.append("cover letter closing is required")
        if not model.signature.strip():
            errors.append("cover letter signature is required")
    return {"ok": not errors, "errors": errors, "doc_type": model.doc_type}


def check_content_integrity(
    source_texts: Sequence[str],
    extracted_text: str,
    *,
    footer_name: str = "",
) -> Dict[str, Any]:
    """Every source token must survive into the rendered document.

    Returns {ok, missing_tokens, markup_leaks, source_tokens, extracted_tokens}.
    """
    source: Dict[str, int] = {}
    for text in source_texts:
        for token in _tokens(text or ""):
            source[token] = source.get(token, 0) + 1

    extracted: Dict[str, int] = {}
    for token in _tokens(extracted_text or ""):
        extracted[token] = extracted.get(token, 0) + 1

    missing = sorted(token for token, count in source.items() if extracted.get(token, 0) < count)

    # markup leak check on the extracted text (footer name is legit text)
    scan_text = extracted_text or ""
    if footer_name:
        scan_text = scan_text.replace(footer_name, "")
    leaks = sorted({match.group(0) for match in _MARKUP_LEAK.finditer(scan_text)})

    return {
        "ok": not missing and not leaks,
        "missing_tokens": missing,
        "markup_leaks": leaks,
        "source_tokens": len(source),
        "extracted_tokens": len(extracted),
    }


def model_source_texts(model: DocumentModel) -> List[str]:
    """Every content string the model approved for rendering (structured source).

    This is the integrity reference: the renderer may add layout artifacts, but
    it must not lose a single approved word. Collected from the model itself so
    the check can never pass because of a stale hand-written list.
    """
    parts: List[str] = [model.candidate.name, model.candidate.professional_title,
                        model.candidate.status_line]
    for item in model.candidate.contact_items:
        parts.append(item.text if isinstance(item, Link) else str(item))
    for section in model.sections:
        parts.append(section.heading)
        parts.extend(paragraph.text for paragraph in section.paragraphs)
        for bullet in section.bullets:
            parts.extend([bullet.label, bullet.text])
        for entry in section.experience:
            parts.extend([entry.role, entry.company, entry.location, entry.period])
            parts.extend(paragraph.text for paragraph in entry.paragraphs)
            parts.extend([bullet.label + " " + bullet.text for bullet in entry.bullets])
        for entry in section.education:
            parts.extend([entry.heading, entry.metadata])
        for entry in section.entries:
            parts.extend([entry.text, entry.metadata])
    # cover letter specifics
    parts.append(model.contact_block_lead)
    for item in model.contact_block:
        parts.append(item.text if isinstance(item, Link) else str(item))
    parts.extend([model.date_line, model.location_line, model.subject,
                  model.salutation, model.closing, model.signature])
    return [part for part in parts if part and part.strip()]


def validate_rendered_document(
    model: DocumentModel,
    docx_bytes: bytes,
    source_texts: Optional[Sequence[str]] = None,
) -> Dict[str, Any]:
    """Full DOCX QA round: schema + extraction + integrity in one report."""
    schema = validate_document_model(model)
    extracted = _extract_docx_text(docx_bytes)
    sources = list(source_texts) if source_texts else model_source_texts(model)
    integrity = check_content_integrity(sources, extracted, footer_name=model.candidate.name)
    return {
        "doc_type": model.doc_type,
        "schema": schema,
        "integrity": integrity,
        "ok": schema["ok"] and integrity["ok"],
        "errors": list(schema["errors"]) + (
            [f"missing content: {', '.join(integrity['missing_tokens'][:10])}"]
            if integrity["missing_tokens"] else []
        ) + (
            [f"markup leak: {', '.join(integrity['markup_leaks'][:5])}"]
            if integrity["markup_leaks"] else []
        ),
    }
